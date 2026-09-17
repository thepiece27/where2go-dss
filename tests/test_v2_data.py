from datetime import date
import sqlite3
from openpyxl import load_workbook

from scripts.create_manual_template_v2 import create
from scripts.validate_manual_data_v2 import validate
from where2go.config import CATALOG
from where2go.v2.durations import choose_duration, fallback_profile
from where2go.v2.hours import intervals_on_date
from where2go.v2.storage import connect, create_database
from where2go.v2.taxonomy import canonical_category, refine_category
from where2go.v2.quality import serving_quality


def test_taxonomy_specific_rules_precede_broad_park():
    assert canonical_category("National park") == "nature_area"
    assert canonical_category("Coffee shop") == "cafe"
    assert canonical_category("Results") is None
    assert refine_category("attraction", "Phố cổ Hội An") == "old_quarter"
    assert refine_category("attraction", "Làng gốm Thanh Hà") == "craft_village"
    assert refine_category("attraction", "Nhà hát múa rối") == "performance_venue"


def test_duration_profile_and_override_are_explicit():
    profile = fallback_profile("museum")
    assert choose_duration(profile, "quick") == (30, "category_default")
    assert choose_duration(profile, "relaxed") == (120, "category_default")
    assert choose_duration(profile, "balanced", 75) == (75, "user_override")


def test_hours_unknown_closed_and_overnight_are_distinct():
    weekly = [None, [], [(1320, 1560)], [], [], [], []]
    assert intervals_on_date(weekly, {}, date(2026, 9, 14), True) is None  # Monday unknown
    assert intervals_on_date(weekly, {}, date(2026, 9, 15), True) == []   # Tuesday closed
    assert intervals_on_date(weekly, {}, date(2026, 9, 17), True) == [(0, 120)]
    exceptions = {"2026-09-17": {"status": "closed"}}
    assert intervals_on_date(weekly, exceptions, date(2026, 9, 17), True) == []


def test_storage_schema_enforces_foreign_keys(tmp_path):
    path = tmp_path / "v2.sqlite"
    create_database(path)
    with connect(path) as db:
        assert db.execute("SELECT value FROM metadata WHERE key='schema_version'").fetchone()[0] == "2.0"
        try:
            db.execute("INSERT INTO source_records VALUES ('r','missing','1',NULL,'{}')")
        except sqlite3.IntegrityError:
            pass
        else:
            raise AssertionError("Foreign key was not enforced")


def test_manual_template_is_valid_before_human_enrichment(tmp_path):
    path = tmp_path / "manual.xlsx"
    count = create(path, CATALOG, per_city=2)
    errors, warnings, stats = validate(path)
    assert count == 4
    assert not errors
    assert not warnings
    assert stats["places"] == 4
    workbook = load_workbook(path, read_only=True)
    assert {"places", "opening_hours", "visit_duration", "verification"} <= set(workbook.sheetnames)


def test_serving_quality_blocks_generic_and_evidence_free_pois():
    base = {"name": "Hard to climb", "ratings": [], "hours_weekly": None,
            "website": "", "description": "", "duration_profile": {"method": "category_default"},
            "access_points": [{"verified": False}]}
    quality = serving_quality(base)
    assert not quality["eligible"]
    assert set(quality["reasons"]) == {"weak_or_generic_name", "insufficient_service_evidence"}
    improved = dict(base, name="Bảo tàng kiểm thử", website="https://example.test")
    assert serving_quality(improved)["eligible"]
    improved["hours_weekly"] = [[(480, 1020)]] * 7
    assert serving_quality(improved)["eligible"]


def test_quality_detects_name_category_conflict_and_repeated_description():
    row = {"name": "Trại gà Đông Tảo", "category": "historic", "ratings": [], "hours_weekly": None,
           "website": "2lualamgiau.com", "description": "Trại gà Đông Tảo",
           "duration_profile": {"method": "category_default"}, "access_points": [{"verified": False}]}
    quality = serving_quality(row)
    assert not quality["eligible"]
    assert "category_name_conflict" in quality["reasons"]
    assert not quality["components"]["description"]
