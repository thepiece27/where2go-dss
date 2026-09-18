"""Field-level national source reconciliation; ambiguous observations stay raw."""
from collections import Counter, defaultdict
import json
from pathlib import Path
import re
from urllib.parse import urlparse

from openpyxl import load_workbook
from where2go.catalog import haversine
from where2go.config import ROOT
from where2go.ranking import normalize
from where2go.v2.google_collector import google_identity, image_url_allowed, name_similarity
from where2go.v2.observations import entity_coordinate, parse_rating, parse_review_count, scalar, stable_id
from where2go.v2.taxonomy import canonical_category, tags_for
from where2go.v2.durations import fallback_profile
from where2go.v2.quality import weak_name
from where2go.v2.google_hours import parse_google_week
from scripts.build_catalog import Boundaries


def empty(value):
    return value is None or (isinstance(value, str) and value.strip().lower() in ("", "nan", "none", "null"))


def rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook:
            iterator = sheet.iter_rows(values_only=True)
            headers = next(iterator)
            for number, values in enumerate(iterator, 2):
                row = {str(k).strip(): None if empty(v) else scalar(v) for k, v in zip(headers, values)}
                if any(v is not None for v in row.values()):
                    yield sheet.title, number, row
    finally:
        workbook.close()


def current_pois(db):
    pois = {r["poi_id"]: {"poi_id": r["poi_id"], "name": r["display_name"], "location": r["location"],
                            "data_status": r["status"]} for r in db.execute("SELECT * FROM pois")}
    for r in db.execute("SELECT poi_id,field_name,value_json FROM selected_fields"):
        pois[r["poi_id"]][r["field_name"]] = json.loads(r["value_json"])
    for r in db.execute("SELECT poi_id,name FROM poi_aliases"):
        pois[r["poi_id"]].setdefault("aliases", []).append(r["name"])
    return pois


def compatible(a, b):
    return not a or not b or a == b or "attraction" in (a, b)


class Matcher:
    def __init__(self, pois, geometries=None):
        self.pois = pois
        self.names = defaultdict(set)
        self.ids = {}
        self.geometries = geometries or {}
        self.cells = defaultdict(set)
        for p in pois.values():
            self.index(p)

    def index(self, p):
        for name in [p["name"], *p.get("aliases", [])]:
            self.names[normalize(name)].add(p["poi_id"])
        self.cells[(int(p["latitude"] * 100), int(p["longitude"] * 100))].add(p["poi_id"])

    def possible_duplicates(self, name, coordinate, category):
        y, x = (int(v * 100) for v in coordinate)
        found = []
        for dy in (-1, 0, 1):
            for dx in (-1, 0, 1):
                for ident in self.cells.get((y + dy, x + dx), []):
                    p = self.pois[ident]
                    if (compatible(category, p.get("category")) and haversine(coordinate, (p["latitude"], p["longitude"])) <= .3
                            and max(name_similarity(name, n) for n in [p["name"], *p.get("aliases", [])]) >= .6):
                        found.append(ident)
        return sorted(set(found))

    def match(self, names, coordinate, category=None, external_id=None):
        if not coordinate:
            return [], "missing_entity_coordinate"
        candidates = set()
        for name in names:
            candidates.update(self.names.get(normalize(name), set()))
        if external_id in self.ids:
            candidates.add(self.ids[external_id])
        accepted = []
        from shapely.geometry import Point, shape
        for ident in sorted(candidates):
            p = self.pois[ident]
            if p["data_status"] == "excluded" or not compatible(category, p.get("category")):
                continue
            near = haversine(coordinate, (p["latitude"], p["longitude"])) <= .3
            area = self.geometries.get(ident)
            in_area = area and shape(area).covers(Point(coordinate[1], coordinate[0]))
            if near or in_area:
                accepted.append(p)
        return accepted, "name_category_and_entity_position"


def alias(db, poi_id, name, observation):
    if name and not weak_name(name):
        db.execute("INSERT OR IGNORE INTO poi_aliases VALUES (?,?,?,?)", (poi_id, name.strip(), normalize(name), observation))


def select_if_empty(db, poi_id, field, value, observation, version, reason):
    from scripts.build_catalog_v2 import select
    existing = db.execute("SELECT value_json FROM selected_fields WHERE poi_id=? AND field_name=?", (poi_id, field)).fetchone()
    if not empty(value) and (not existing or empty(json.loads(existing[0]))):
        select(db, poi_id, field, observation, value, reason, version)
        return True
    return False


def add_image(db, poi_id, url, source_url, record_id, observed_at, identity_status):
    from scripts.build_catalog_v2 import observe
    if not isinstance(url, str) or not url:
        return False
    observation = observe(db, poi_id, record_id, "image_candidate", url, observed_at,
                          "restricted_internal", "source_image_candidate", notes=identity_status)
    provider = "Google Maps" if "google" in url else "legacy website"
    state = "pending" if image_url_allowed(url) else "rejected_url"
    db.execute("INSERT OR IGNORE INTO poi_images VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)", (
        stable_id("image", poi_id, url), poi_id, url, source_url, provider, "restricted_internal",
        identity_status, state, None, None, None, None, observation))
    return state == "pending"


def legacy_image_identity(url, result, category):
    """An arbitrary website image in the old row is not evidence of a Maps panel."""
    if weak_name(result) or not category or not image_url_allowed(url):
        return "needs_review"
    parsed = urlparse(url)
    google_photo = ((parsed.hostname or "").endswith(".googleusercontent.com")
                    and parsed.path.startswith(("/gps-cs-s/", "/gps-proxy/", "/p/", "/places/")))
    if google_photo or parsed.hostname == "streetviewpixels-pa.googleapis.com":
        return "legacy_entity_link"
    return "needs_review"


def import_legacy(db, paths, queue, version, geometries=None):
    from scripts.build_catalog_v2 import add_source, add_record, observe, select
    matcher = Matcher(current_pois(db), geometries)
    for r in db.execute("SELECT external_id,poi_id FROM external_ids WHERE provider='Google Maps'"):
        matcher.ids[r[0]] = r[1]
    boundaries = Boundaries(ROOT / "data/vietnam_provinces_wards_geojson.zip")
    stats = Counter()
    source_stats = {}
    linked_pois = set()
    ids_seen = set()
    rating_seen = set()
    name_location_evidence = defaultdict(set)
    snapshots = []
    for path in paths:
        snapshot = list(rows(path))
        snapshots.append((path, snapshot))
        for _, _, row in snapshot:
            location = normalize(row.get("Vị trí") or "")
            if location in {normalize(name) for name, _ in boundaries.items}:
                name_location_evidence[str(row.get("STT"))].add(location)
    for path, snapshot in snapshots:
        source_id = add_source(db, path, "source_observation_restricted", "restricted_internal", "legacy Excel reconciliation")
        counts = Counter()
        for sheet, number, row in snapshot:
            key = str(row.get("STT") or f"row:{number}")
            ids_seen.add(key)
            record = add_record(db, source_id, sheet + ":" + key, {**row, "_sheet": sheet, "_row": number}, None)
            counts["rows"] += 1
            seed = row.get("Tên địa điểm") or ""
            result = row.get("maps_result_name") or ""
            url = row.get("maps_url") or ""
            coordinate = entity_coordinate(url)
            external = google_identity(url)
            category = canonical_category(row.get("maps_destination_type"), result)
            candidates, method = matcher.match([seed, result], coordinate, category, external)
            reasons = []
            location = boundaries.locate(coordinate[1], coordinate[0]) if coordinate else ""
            locations = name_location_evidence[key]
            if len(locations) > 1 or (location and locations and normalize(location) not in locations):
                reasons.append("conflicting_snapshot_location")
            if not external:
                reasons.append("missing_external_identity")
            if not coordinate:
                reasons.append("missing_entity_coordinate")
            if len(candidates) != 1:
                reasons.append("multiple_candidates" if candidates else "no_conservative_match")
            if weak_name(seed):
                reasons.append("weak_seed_name")
            can_create = (not candidates and coordinate and external and external not in matcher.ids
                          and category and location and len(locations) == 1 and normalize(location) in locations
                          and normalize(seed) == normalize(result) and not weak_name(seed)
                          and row.get("maps_match_status") == "matched")
            if can_create:
                possible = matcher.possible_duplicates(seed, coordinate, category)
                if not possible:
                    ident = stable_id("poi-google", external)
                    insert_poi(db, ident, seed, location, category, coordinate, record, version, url)
                    poi = {"poi_id": ident, "name": seed, "location": location, "category": category,
                           "latitude": coordinate[0], "longitude": coordinate[1], "data_status": "usable"}
                    matcher.pois[ident] = poi
                    matcher.index(poi)
                    candidates, reasons = [poi], []
                    counts["new_pois"] += 1
                    method = "legacy_exact_entity_name_category_and_boundary"
                else:
                    reasons.append("nearby_similar_entity_requires_review")
                    candidates = [matcher.pois[ident] for ident in possible]
            if reasons:
                counts["ambiguous" if len(candidates) > 1 else "unmatched"] += 1
                queue.append({"source_file": path.name, "source_key": key, "sheet": sheet, "row": number,
                              "seed_name": seed, "result_name": result, "status": "unmatched",
                              "candidate_poi_ids": [p["poi_id"] for p in candidates], "reasons": reasons})
                continue
            poi = candidates[0]
            ident = poi["poi_id"]
            if external in matcher.ids and matcher.ids[external] != ident:
                counts["external_identity_conflict"] += 1
                queue.append({"source_file": path.name, "source_key": key, "seed_name": seed,
                              "candidate_poi_ids": [ident, matcher.ids[external]], "reasons": ["external_identity_conflict"]})
                continue
            db.execute("INSERT INTO source_links VALUES (?,?,?,?,?,?,?)", (ident, record, method, 1., None, "confirmed", "Legacy rating time remains unknown"))
            counts["confirmed"] += 1
            linked_pois.add(ident)
            if not poi.get("location") and location:
                observation = observe(db, ident, record, "location", location, None, "open_data_derived",
                                      "cross_source_spatial_assignment", notes="Matched entity marker is covered by the province boundary")
                select(db, ident, "location", observation, location, "OSM identity plus entity-coordinate boundary assignment", version)
                db.execute("UPDATE pois SET location=? WHERE poi_id=?", (location, ident))
                poi["location"] = location
                counts["spatial_assignments"] += 1
            obs = observe(db, ident, record, "google_maps_url", url, None, "restricted_internal", method)
            db.execute("INSERT OR IGNORE INTO external_ids VALUES ('Google Maps',?,?,?)", (external, ident, obs))
            matcher.ids[external] = ident
            alias(db, ident, seed, obs)
            # The result name belongs to the URL entity only when it is a real panel name.
            if not weak_name(result) and category and compatible(category, poi.get("category")):
                alias(db, ident, result, obs)
            for field, value in (("description", row.get("Mô tả")), ("keywords", row.get("Từ Khóa")),
                                 ("address_raw", row.get("Vị trí")), ("google_category_raw", row.get("maps_destination_type")),
                                 ("google_hours_raw", row.get("maps_open_hours")), ("google_maps_url", url)):
                if not empty(value):
                    observation = observe(db, ident, record, field, value, None, "restricted_internal", "legacy_field_unverified")
                    if select_if_empty(db, ident, field, value, observation, version, "Recovered nonempty legacy field; observation time unknown"):
                        counts["selected_fields_recovered"] += 1
            rating, count = parse_rating(row.get("Đánh giá")), parse_review_count(row.get("maps_review_count"))
            signature = (ident, rating, count)
            if (rating is not None or count is not None) and signature not in rating_seen:
                observation = observe(db, ident, record, "legacy_rating", {"rating": rating, "review_count": count}, None,
                                      "restricted_internal", "legacy_pair_time_unproven")
                db.execute("INSERT INTO ratings VALUES (?,?,?,?,?,?,?,?)", (stable_id("legacy-rating", *signature), ident, "Google Maps (legacy unverified)", rating, count, None, 0, observation))
                rating_seen.add(signature)
            image_identity = legacy_image_identity(row.get("Ảnh"), result, category)
            if add_image(db, ident, row.get("Ảnh"), url, record, None, image_identity):
                counts["image_candidates"] += 1
        source_stats[path.name] = dict(counts)
        stats.update(counts)
    return {**stats, "unique_legacy_ids": len(ids_seen), "linked_unique_pois": len(linked_pois), "sources": source_stats}


def insert_poi(db, ident, name, location, category, coordinate, record, version, source_url):
    from scripts.build_catalog_v2 import observe, select
    db.execute("INSERT INTO pois VALUES (?,?,?,?,?,?)", (ident, "usable", name, location, None, "unknown"))
    for field, value in {"name": name, "location": location, "category": category, "latitude": coordinate[0],
                         "longitude": coordinate[1], "source_url": source_url}.items():
        obs = observe(db, ident, record, field, value, None, "restricted_internal", "tool_cross_source_identity")
        select(db, ident, field, obs, value, "New entity with source identity and spatial evidence", version)
    db.execute("INSERT INTO poi_categories VALUES (?,?,1,NULL)", (ident, category))
    for tag in tags_for(category):
        db.execute("INSERT INTO poi_categories VALUES (?,?,0,NULL)", (ident, "tag:" + tag))
    profile = fallback_profile(category)
    db.execute("INSERT INTO duration_profiles VALUES (?,?,?,?,?,NULL,NULL)", (ident, profile["short_minutes"], profile["typical_minutes"], profile["long_minutes"], profile["method"]))
    db.execute("INSERT INTO access_points VALUES (?,?,?,?,?,?,0,?,NULL)", (stable_id("access", ident), ident, *coordinate, "poi_coordinate", 10, "google_entity_not_verified_entrance"))


def import_collected(db, paths, queue, version):
    from scripts.build_catalog_v2 import add_source, add_record, observe, select
    counts = Counter()
    for path in paths:
        if not path.exists():
            continue
        source = add_source(db, path, "google_collection", "restricted_internal", "where2go collector adapter")
        attempts = json.loads(path.read_text(encoding="utf-8-sig"))
        latest = {r["record_id"]: r for r in attempts}
        for row in latest.values():
            record = add_record(db, source, row["record_id"], row, row.get("scraped_at"))
            counts["attempts"] += 1
            place = row.get("place") or {}
            if (row.get("identity") or {}).get("status") != "tool_confirmed" or row.get("accepted") is not True:
                queue.append({"source_file": path.name, "source_key": row["record_id"], "seed_name": row.get("seed_name"),
                              "reasons": ["collection_identity_review_required"], "candidate_poi_ids": [row.get("canonical_poi_id")]})
                continue
            url = place.get("google_maps_url") or place.get("url")
            coordinate = entity_coordinate(url)
            external = google_identity(url)
            if not coordinate or not external:
                continue
            linked = db.execute("SELECT poi_id FROM external_ids WHERE provider='Google Maps' AND external_id=?", (external,)).fetchone()
            ident = row.get("canonical_poi_id") or (linked[0] if linked else stable_id("poi-google", external))
            if linked and linked[0] != ident:
                counts["external_identity_conflict"] += 1
                queue.append({"source_file": path.name, "source_key": row["record_id"], "seed_name": row.get("seed_name"),
                              "candidate_poi_ids": [ident, linked[0]], "reasons": ["external_identity_conflict"]})
                continue
            if not db.execute("SELECT 1 FROM pois WHERE poi_id=?", (ident,)).fetchone():
                category = canonical_category(place.get("category"), place.get("name")) or row.get("expected_category")
                if not category:
                    continue
                insert_poi(db, ident, row["seed_name"], row["location_expected"], category, coordinate, record, version, url)
                counts["new_pois"] += 1
            db.execute("INSERT OR IGNORE INTO source_links VALUES (?,?,?,?,?,?,?)", (ident, record, "tool_cross_source_identity", 1., "tool:source_review", "confirmed", row.get("review_evidence")))
            timestamp = row.get("scraped_at")
            position_obs = observe(db, ident, record, "entity_coordinate", {"latitude": coordinate[0], "longitude": coordinate[1]}, timestamp,
                                   "restricted_internal", "google_entity_url_not_verified_entrance")
            db.execute("INSERT OR REPLACE INTO access_points VALUES (?,?,?,?,?,?,0,?,?)", (
                stable_id("access", ident, "live-google-entity"), ident, *coordinate, "poi_coordinate", 10,
                "google_entity_url_not_verified_entrance", position_obs))
            id_obs = observe(db, ident, record, "google_maps_url", url, timestamp, "restricted_internal", "live_entity_panel")
            db.execute("INSERT OR IGNORE INTO external_ids VALUES ('Google Maps',?,?,?)", (external, ident, id_obs))
            for name in (row.get("seed_name"), place.get("name")):
                alias(db, ident, name, id_obs)
            if row["record_id"].startswith("focus-"):
                name_obs = observe(db, ident, record, "name", row["seed_name"], None, "internal_design_data", "curated_name_on_confirmed_entity")
                select(db, ident, "name", name_obs, row["seed_name"], "Curated landmark name; original names retained as aliases", version)
                db.execute("UPDATE pois SET display_name=? WHERE poi_id=?", (row["seed_name"], ident))
            for field, value in (("google_maps_url", url), ("address_raw", place.get("address")), ("website", place.get("website")), ("phone", place.get("phone"))):
                if not empty(value):
                    obs = observe(db, ident, record, field, value, timestamp, "restricted_internal", "live_entity_panel")
                    select_if_empty(db, ident, field, value, obs, version, "Accepted live entity panel")
            rating, count = parse_rating(place.get("rating")), parse_review_count(place.get("review_count"))
            if rating is not None and count is not None:
                obs = observe(db, ident, record, "rating_pair", {"rating": rating, "review_count": count}, timestamp, "restricted_internal", "same_live_entity_panel")
                db.execute("INSERT OR REPLACE INTO ratings VALUES (?,?,?,?,?,?,1,?)", (stable_id("rating", record), ident, "Google Maps", rating, count, timestamp, obs))
                counts["rating_pairs"] += 1
            weekly = parse_google_week(place.get("hours"))
            if weekly is not None:
                obs = observe(db, ident, record, "opening_hours", weekly, timestamp, "restricted_internal", "live_entity_weekly_hours")
                # Replace only the days actually observed; unknown days never erase known intervals.
                for day, intervals in enumerate(weekly):
                    if intervals is None:
                        continue
                    db.execute("DELETE FROM opening_intervals WHERE poi_id=? AND day_of_week=? AND specific_date IS NULL", (ident, day))
                    for i, (start, end) in enumerate(intervals or [(None, None)]):
                        db.execute("INSERT INTO opening_intervals VALUES (?,?,?,?,?,?,?,?,?)", (stable_id("hours", record, day, i), ident, obs, day, None, "closed" if start is None else "open", start, end if end is None or end <= 1440 else end - 1440, int(end is not None and end > 1440)))
                counts["structured_hours"] += 1
            if add_image(db, ident, place.get("image_url"), url, record, timestamp, "entity_panel"):
                counts["image_candidates"] += 1
            business = "permanently_closed" if place.get("permanently_closed") else "temporarily_closed" if place.get("temporarily_closed") else None
            if business:
                db.execute("UPDATE pois SET business_status=? WHERE poi_id=?", (business, ident))
            counts["accepted"] += 1
    return dict(counts)


def apply_images(db, cache_path, version):
    from scripts.build_catalog_v2 import select
    checks = json.loads(cache_path.read_text(encoding="utf-8")) if cache_path.exists() else {}
    for url, result in checks.items():
        db.execute("UPDATE poi_images SET validation_status=?,checked_at=?,content_type=?,width=?,height=? WHERE url=? AND validation_status<>'rejected_url'", (
            result["status"], result.get("checked_at"), result.get("content_type"), result.get("width"), result.get("height"), url))
    selected = set()
    # A URL repeated across distinct entities needs review instead of being a representative photo.
    shared = {r[0] for r in db.execute("SELECT url FROM poi_images GROUP BY url HAVING count(DISTINCT poi_id)>1")}
    candidates = db.execute("SELECT i.*,o.observed_at FROM poi_images i LEFT JOIN field_observations o USING(observation_id) WHERE validation_status='valid' AND identity_status IN ('entity_panel','legacy_entity_link') ORDER BY o.observed_at DESC,i.image_id").fetchall()
    for row in candidates:
        if row["poi_id"] in selected or row["url"] in shared:
            continue
        select(db, row["poi_id"], "image", row["observation_id"], row["url"], "Entity-linked image; live content and dimensions checked", version)
        select(db, row["poi_id"], "image_metadata", row["observation_id"], {
            "provider": row["provider"], "source_url": row["source_url"], "rights_status": row["rights_status"],
            "validation_status": row["validation_status"], "checked_at": row["checked_at"],
            "width": row["width"], "height": row["height"], "identity_status": row["identity_status"]}, "Image provenance", version)
        selected.add(row["poi_id"])
    return {"selected": len(selected), "shared_urls_held_for_review": len(shared), "cached_checks": len(checks)}


def reconcile_status(db):
    """Clear only the resolved spatial reason; never clear unrelated review flags."""
    count = 0
    for poi in current_pois(db).values():
        reasons = poi.get("review_reasons", [])
        if poi["data_status"] != "needs_review" or not poi.get("location"):
            continue
        assignment = db.execute("SELECT 1 FROM field_observations WHERE poi_id=? AND field_name='location' AND verification_method IN ('curated_spatial_assignment','cross_source_spatial_assignment')", (poi["poi_id"],)).fetchone()
        if assignment and reasons and set(reasons) <= {"outside_boundary_needs_review"}:
            db.execute("UPDATE pois SET status='usable' WHERE poi_id=?", (poi["poi_id"],))
            db.execute("UPDATE selected_fields SET value_json='[]' WHERE poi_id=? AND field_name='review_reasons'", (poi["poi_id"],))
            count += 1
    return count
