"""Copy successful pilot observations into the manual workbook as pending review."""
import argparse
import csv
from datetime import time
import json
from pathlib import Path
import re
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from scripts.create_manual_template_v2 import PLACES, OPENING, VERIFICATION
from where2go.config import ROOT
from where2go.v2.google_hours import clean_google_text, parse_google_week


MARKER = "collector_observation_pending_human_verification"
PLACE_COORDINATES = re.compile(r"!3d(-?\d+(?:\.\d+)?)!4d(-?\d+(?:\.\d+)?)")


def header_map(sheet):
    return {cell.value: cell.column for cell in sheet[1]}


def as_time(minutes):
    minutes %= 1440
    return time(minutes // 60, minutes % 60)


def import_results(workbook_path, results_path, review_path):
    results = {row["record_id"]: row for row in json.loads(results_path.read_text(encoding="utf-8"))}
    with review_path.open(encoding="utf-8-sig", newline="") as handle:
        review = {row["record_id"]: row for row in csv.DictReader(handle)}
    workbook = load_workbook(workbook_path)
    places = workbook["places"]
    columns = header_map(places)
    marker = f"{MARKER}:{results_path.parent.name}/{results_path.name}"
    existing_places = {
        str(places.cell(row, columns["record_id"]).value or "").strip()
        for row in range(2, places.max_row + 1)
    }
    for record_id, result in results.items():
        if record_id in existing_places:
            continue
        seed = {key: "" for key in PLACES}
        seed.update({
            "record_id": record_id,
            "canonical_poi_id": result.get("canonical_poi_id") or "",
            "seed_name": result.get("seed_name") or "",
            "location_expected": result.get("location_expected") or "",
            "category_raw": "",
            "notes": marker,
        })
        places.append([seed[key] for key in PLACES])
        existing_places.add(record_id)
    collector_fields = (
        "maps_name", "maps_url", "place_id", "cid", "address_raw", "rating_raw",
        "review_count_raw", "latitude", "longitude", "coordinate_method", "website", "category_raw",
        "business_status_raw", "observed_at", "reviewer",
    )
    for row_number in range(2, places.max_row + 1):
        record_id = str(places.cell(row_number, columns["record_id"]).value or "").strip()
        reviewer = str(places.cell(row_number, columns["reviewer"]).value or "").strip()
        if record_id in results and reviewer == "collector:noworneverev":
            for field in collector_fields:
                places.cell(row_number, columns[field]).value = None
    imported = 0
    for row_number in range(2, places.max_row + 1):
        record_id = str(places.cell(row_number, columns["record_id"]).value or "").strip()
        result = results.get(record_id)
        assessment = review.get(record_id, {})
        if not result or not result.get("success") or assessment.get("machine_match", "").lower() != "true":
            continue
        place = result.get("place") or {}
        place_url = place.get("google_maps_url") or place.get("url") or ""
        latitude, longitude = place.get("latitude"), place.get("longitude")
        coordinate_method = "google_place_coordinate" if latitude is not None else ""
        if latitude is None:
            coordinate_match = PLACE_COORDINATES.search(place_url)
            if coordinate_match:
                latitude, longitude = map(float, coordinate_match.groups())
                coordinate_method = "google_place_url_data_coordinate"
        status = "permanently_closed" if place.get("permanently_closed") else (
            "temporarily_closed" if place.get("temporarily_closed") else "open"
        )
        values = {
            "maps_name": clean_google_text(place.get("name")),
            "maps_url": place_url,
            "place_id": place.get("place_id"), "address_raw": clean_google_text(place.get("address")),
            "category_raw": clean_google_text(place.get("category")), "rating_raw": place.get("rating"),
            "review_count_raw": place.get("review_count"), "latitude": latitude,
            "longitude": longitude, "coordinate_method": coordinate_method,
            "website": place.get("website"), "business_status_raw": status,
            "observed_at": result.get("scraped_at"), "reviewer": "collector:noworneverev",
            "notes": marker,
        }
        for field, value in values.items():
            if value not in (None, "") and places.cell(row_number, columns[field]).value in (None, ""):
                places.cell(row_number, columns[field]).value = value
        imported += 1

    opening = workbook["opening_hours"]
    opening_columns = header_map(opening)
    for row_number in range(opening.max_row, 1, -1):
        record_id = str(opening.cell(row_number, opening_columns["record_id"]).value or "").strip()
        notes = str(opening.cell(row_number, opening_columns["notes"]).value or "")
        if record_id in results and notes.startswith(MARKER):
            opening.delete_rows(row_number)
    day_codes = ("Mo", "Tu", "We", "Th", "Fr", "Sa", "Su")
    opening_rows = 0
    for record_id, result in results.items():
        assessment = review.get(record_id, {})
        if not result.get("success") or assessment.get("machine_match", "").lower() != "true":
            continue
        place = result.get("place") or {}
        weekly = parse_google_week(place.get("hours"))
        if weekly is None:
            continue
        source_url = place.get("google_maps_url") or place.get("url") or result.get("input_url")
        for day_index, intervals in enumerate(weekly):
            if intervals is None:
                continue
            if not intervals:
                payload = {"record_id": record_id, "day_of_week": day_codes[day_index],
                           "status": "closed", "source_url": source_url,
                           "observed_at": result.get("scraped_at"), "notes": marker}
                opening.append([payload.get(field, "") for field in OPENING]); opening_rows += 1
                continue
            for start, end in intervals:
                payload = {
                    "record_id": record_id, "day_of_week": day_codes[day_index], "status": "open",
                    "open_time": as_time(start), "close_time": as_time(end),
                    "closes_next_day": end > 1440 or (start == 0 and end == 1440),
                    "source_url": source_url, "observed_at": result.get("scraped_at"), "notes": marker,
                }
                opening.append([payload.get(field, "") for field in OPENING]); opening_rows += 1

    verification = workbook["verification"]
    verification_columns = header_map(verification)
    existing = {
        str(verification.cell(row, verification_columns["record_id"]).value or "").strip()
        for row in range(2, verification.max_row + 1)
    }
    verification_rows = 0
    for record_id, result in results.items():
        if record_id in existing or not result.get("success"):
            continue
        place = result.get("place") or {}
        payload = {
            "record_id": record_id, "reviewer": "collector:noworneverev",
            "reviewed_at": result.get("scraped_at"), "fields_checked": "tool_assisted_only",
            "identity_status": "ambiguous",
            "evidence_url": place.get("google_maps_url") or place.get("url") or result.get("input_url"),
            "include_in_demo": False,
            "notes": "Cần người kiểm duyệt xác nhận đúng thực thể trước khi nhập catalog.",
        }
        verification.append([payload.get(field, "") for field in VERIFICATION]); verification_rows += 1
    workbook.save(workbook_path)
    return {"places_enriched": imported, "opening_rows": opening_rows, "verification_rows": verification_rows}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=ROOT / "data/manual/poi_enrichment_v2.xlsx")
    parser.add_argument("--results", type=Path, default=ROOT / "data/google-pilot/results.json")
    parser.add_argument("--review", type=Path, default=ROOT / "data/google-pilot/review.csv")
    args = parser.parse_args()
    print(json.dumps(import_results(args.workbook, args.results, args.review), ensure_ascii=False))


if __name__ == "__main__":
    main()
