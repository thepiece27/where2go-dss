"""Prepare deterministic gosom/google-maps-scraper inputs from the manual pilot."""
import argparse
import csv
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from where2go.config import ROOT


def prepare(workbook_path, output_dir, limit=None):
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["places"]
    headers = [cell.value for cell in sheet[1]]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        record_id = str(row.get("record_id") or "").strip()
        seed_name = str(row.get("seed_name") or "").strip()
        location = str(row.get("location_expected") or "").strip()
        if record_id and seed_name and location:
            rows.append({
                "record_id": record_id,
                "canonical_poi_id": str(row.get("canonical_poi_id") or "").strip(),
                "seed_name": seed_name,
                "location_expected": location,
                "query": f"{seed_name}, {location}, Việt Nam",
            })
    if limit is not None:
        rows = rows[:limit]
    output_dir.mkdir(parents=True, exist_ok=True)
    query_path = output_dir / "queries.txt"
    query_path.write_text("".join(f"{row['query']} #!#{row['record_id']}\n" for row in rows), encoding="utf-8")
    with (output_dir / "seeds.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]) if rows else [
            "record_id", "canonical_poi_id", "seed_name", "location_expected", "query",
        ])
        writer.writeheader()
        writer.writerows(rows)
    (output_dir / "README.md").write_text(
        "# Google Maps pilot v2\n\n"
        "`queries.txt` là input có ID cho gosom/google-maps-scraper; `seeds.csv` dùng để đối sánh. "
        "Raw result không tự trở thành catalog. Chỉ bản ghi đúng thực thể đã được xác nhận mới được nâng vào catalog. "
        "Nếu công cụ gặp CAPTCHA/chặn truy cập thì dừng; không vượt chặn. Dữ liệu Google Maps giữ trạng thái "
        "`restricted_internal` và không được coi là dataset công khai.\n",
        encoding="utf-8",
    )
    return rows


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=ROOT / "data/manual/poi_enrichment_v2.xlsx")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "data/google-pilot")
    parser.add_argument("--limit", type=int)
    args = parser.parse_args()
    rows = prepare(args.workbook, args.output_dir, args.limit)
    print(f"Prepared {len(rows)} pilot queries in {args.output_dir}")


if __name__ == "__main__":
    main()
