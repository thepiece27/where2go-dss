"""Run 16 v2 scenarios on four ranking methods through the shared service."""
import argparse
from datetime import date
import hashlib
import json
from pathlib import Path
import random
import sys
import time

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from where2go.config import ROOT
from where2go.routing import OSRM
from where2go.v2.catalog import CATALOG_V2, load_catalog
from where2go.v2.models import ItineraryRequestV2
from where2go.v2.service import ItineraryService


METHODS = ("nearby", "equal", "crisp", "fuzzy")


def scenario_request(row):
    payload = dict(row)
    payload.pop("scenario_id")
    payload.pop("split")
    payload.setdefault("date", date(2026, 9, 18).isoformat())
    payload.setdefault("start_time", "08:00")
    payload.setdefault("end_time", "18:00")
    return ItineraryRequestV2.model_validate(payload)


def summarize(result):
    attraction = [block for block in result.get("blocks", []) if block["role"] == "attraction"]
    return {
        "status": result["status"], "reason": result["reason"],
        "attraction_count": len(attraction), "block_count": len(result.get("blocks", [])),
        "selected_poi_ids": [block["poi_id"] for block in attraction],
        "selected_names": [block["name"] for block in attraction],
        "categories": [block["category"] for block in attraction],
        "drive_minutes": round(result.get("drive_seconds", 0) / 60, 2),
        "return_time": result.get("return_time"), "reserve_minutes": result.get("reserve_minutes"),
        "objective": result.get("objective"), "warnings": result.get("warnings", []),
        "candidate_counts": result.get("candidate_counts", {}), "ranking": result.get("ranking"),
    }


def grading_workbook(results, method_key, output):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "grading"
    headers = ["scenario_id", "split", "itinerary_code", "status", "itinerary_summary",
               "preference_fit_1_5", "poi_value_1_5", "diversity_1_5", "duration_1_5",
               "pace_1_5", "meal_break_1_5", "would_use_1_5", "issues", "replacement"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="155E75")
    for row in results:
        code = method_key[row["scenario_id"]][row["method"]]
        summary = " | ".join(f"{name} [{category}]" for name, category in zip(row["selected_names"], row["categories"]))
        sheet.append([row["scenario_id"], row["split"], code, row["status"], summary, "", "", "", "", "", "", "", "", ""])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A":16,"B":14,"C":16,"D":18,"E":90,"M":45,"N":45}.items():
        sheet.column_dimensions[column].width = width
    notes = workbook.create_sheet("instructions")
    notes.append(["Phiếu do chủ dự án chấm; chưa phải đánh giá người dùng độc lập."])
    notes.append(["Chấm 1–5, không xem method_key trước khi hoàn tất. Để trống khi chưa chấm."])
    notes.append(["AP/NDCG không được tính từ phiếu itinerary này nếu chưa có nhãn POI cho toàn candidate pool."])
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def run(args):
    scenarios = json.loads(args.scenarios.read_text(encoding="utf-8"))
    if len(scenarios) != 16 or sum(row["split"] == "development" for row in scenarios) != 4:
        raise ValueError("Evaluator v2 cần đúng 16 kịch bản: 4 development và 12 holdout")
    pois, manifest = load_catalog(args.catalog)
    service = ItineraryService(pois, manifest, OSRM())
    results = []
    for scenario in scenarios:
        request = scenario_request(scenario)
        for method in METHODS:
            started = time.perf_counter()
            result = service.plan(request, method=method)
            results.append({
                "scenario_id": scenario["scenario_id"], "split": scenario["split"], "method": method,
                "runtime_seconds": round(time.perf_counter() - started, 4), **summarize(result),
            })
            print(f"{scenario['scenario_id']} {method}: {result['status']}", flush=True)
    rng = random.Random(20260917)
    method_key = {}
    for scenario in scenarios:
        codes = ["A", "B", "C", "D"]
        rng.shuffle(codes)
        method_key[scenario["scenario_id"]] = dict(zip(METHODS, codes))
    payload = {
        "dataset_version": manifest["version"], "scenario_sha256": hashlib.sha256(args.scenarios.read_bytes()).hexdigest(),
        "methods": list(METHODS), "note": "Case study; owner grading is NOT RUN until worksheet is completed.",
        "results": results,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    args.method_key.write_text(json.dumps(method_key, ensure_ascii=False, indent=2), encoding="utf-8")
    grading_workbook(results, method_key, args.grading)
    statuses = {status: sum(row["status"] == status for row in results) for status in sorted({row["status"] for row in results})}
    print(json.dumps({"runs": len(results), "statuses": statuses, "output": str(args.output)}, ensure_ascii=False))


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--catalog", type=Path, default=CATALOG_V2)
    parser.add_argument("--scenarios", type=Path, default=ROOT / "data/evaluation/v2_scenarios.json")
    parser.add_argument("--output", type=Path, default=ROOT / "data/reports/v2/evaluation.json")
    parser.add_argument("--method-key", type=Path, default=ROOT / "data/reports/v2/evaluation_method_key.json")
    parser.add_argument("--grading", type=Path, default=ROOT / "data/manual/v2_owner_grading.xlsx")
    args = parser.parse_args()
    run(args)


if __name__ == "__main__":
    main()

