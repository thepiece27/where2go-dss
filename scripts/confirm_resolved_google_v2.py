"""Promote only deterministic place-URL plus cross-source matches to tool_confirmed."""
import argparse
import csv
from datetime import datetime, timezone
import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from where2go.config import ROOT


def columns(sheet):
    return {cell.value: cell.column for cell in sheet[1]}


def confirm(workbook_path, results_path, review_path):
    results = {row["record_id"]: row for row in json.loads(results_path.read_text(encoding="utf-8"))}
    with review_path.open(encoding="utf-8-sig", newline="") as handle:
        review = {row["record_id"]: row for row in csv.DictReader(handle)}
    workbook = load_workbook(workbook_path)
    sheet = workbook["verification"]
    fields = columns(sheet)
    changed = 0
    for row_number in range(2, sheet.max_row + 1):
        record_id = str(sheet.cell(row_number, fields["record_id"]).value or "").strip()
        result = results.get(record_id)
        assessment = review.get(record_id, {})
        place = (result or {}).get("place") or {}
        eligible = (
            assessment.get("machine_match", "").lower() == "true"
            and assessment.get("resolved_place", "").lower() == "true"
            and assessment.get("trusted_resolution", "").lower() == "true"
            and bool(place.get("place_id"))
        )
        if not eligible:
            continue
        evidence = place.get("google_maps_url") or place.get("url") or result.get("input_url")
        values = {
            "reviewer": "tool:resolve_google_seeds_v2",
            "reviewed_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
            "fields_checked": "identity:name,place_url,place_id,location_or_osm_proximity",
            "identity_status": "tool_confirmed", "evidence_url": evidence,
            "conflicts": "", "include_in_demo": True,
            "notes": "Xác nhận chéo bằng URL place cụ thể và tên + địa phương hoặc khoảng cách OSM; chưa kiểm duyệt thủ công.",
        }
        for field, value in values.items():
            sheet.cell(row_number, fields[field]).value = value
        changed += 1
    workbook.save(workbook_path)
    return changed


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=ROOT / "data/manual/poi_enrichment_v2.xlsx")
    parser.add_argument("--results", type=Path, default=ROOT / "data/google-focus-resolved/results.json")
    parser.add_argument("--review", type=Path, default=ROOT / "data/google-focus-resolved/review.csv")
    args = parser.parse_args()
    print(json.dumps({"tool_confirmed": confirm(args.workbook, args.results, args.review)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
