"""Export the v2 SQLite catalog into human-readable CSV and JSON artifacts."""
import argparse
import csv
from contextlib import closing
import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from where2go.config import ROOT
from where2go.v2.catalog import CATALOG_V2, load_catalog
from where2go.v2.dataset import dataset_summary, json_cell
from where2go.v2.quality import latest_rating_pair
from where2go.v2.storage import connect
from openpyxl import load_workbook


def write_csv(path, fieldnames, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def export_private_enrichment(workbook_path, private_dir):
    if not workbook_path.exists():
        return {}
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    verification = {}
    if "verification" in workbook.sheetnames:
        sheet = workbook["verification"]
        headers = [cell.value for cell in sheet[1]]
        for values in sheet.iter_rows(min_row=2, values_only=True):
            row = dict(zip(headers, values))
            if row.get("record_id"):
                verification[str(row["record_id"])] = row
    private_dir.mkdir(parents=True, exist_ok=True)
    sheet = workbook["places"]
    headers = [cell.value for cell in sheet[1]]
    rows = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        if not row.get("maps_name") and not row.get("maps_url"):
            continue
        decision = verification.get(str(row.get("record_id")), {})
        rows.append({
            **{field: row.get(field) or "" for field in (
                "record_id", "canonical_poi_id", "seed_name", "maps_name", "maps_url", "place_id", "cid",
                "address_raw", "location_expected", "category_raw", "rating_raw", "review_count_raw",
                "latitude", "longitude", "coordinate_method", "website", "business_status_raw",
                "observed_at", "reviewer", "notes",
            )},
            "identity_status": decision.get("identity_status") or "unreviewed",
            "human_reviewer": decision.get("reviewer") or "",
            "reviewed_at": decision.get("reviewed_at") or "",
        })
    place_path = private_dir / "google_enrichment_v2.csv"
    write_csv(place_path, list(rows[0]) if rows else ["record_id"], rows)

    sheet = workbook["opening_hours"]
    headers = [cell.value for cell in sheet[1]]
    hours = [dict(zip(headers, values)) for values in sheet.iter_rows(min_row=2, values_only=True)
             if any(value not in (None, "") for value in values)]
    hours_path = private_dir / "google_opening_hours_v2.csv"
    write_csv(hours_path, headers, hours)
    return {
        "google_enrichment": str(place_path.relative_to(ROOT)).replace("\\", "/"),
        "google_opening_hours": str(hours_path.relative_to(ROOT)).replace("\\", "/"),
        "enrichment_rows": len(rows), "opening_rows": len(hours),
    }


def export(catalog, output_dir, workbook_path=None, private_dir=None):
    pois, manifest = load_catalog(catalog)
    output_dir.mkdir(parents=True, exist_ok=True)
    flat = []
    for poi in sorted(pois, key=lambda item: item["poi_id"]):
        rating = latest_rating_pair(poi)
        profile = poi.get("duration_profile") or {}
        access = next(iter(poi.get("access_points") or []), {})
        weekly = poi.get("hours_weekly")
        flat.append({
            "poi_id": poi["poi_id"], "status": poi.get("data_status"), "name": poi.get("name"),
            "location": poi.get("location"), "category": poi.get("category"),
            "tags": "|".join(poi.get("tags") or []), "parent_poi_id": poi.get("parent_poi_id") or "",
            "business_status": poi.get("business_status"), "latitude": poi.get("latitude"),
            "longitude": poi.get("longitude"), "website": poi.get("website") or "",
            "source_url": poi.get("source_url") or "", "description": poi.get("description") or "",
            "hours_raw": poi.get("hours_raw") or "", "has_structured_hours": weekly is not None,
            "has_full_week_hours": weekly is not None and all(day is not None for day in weekly),
            "opening_hours_json": json_cell(weekly),
            "rating": rating.get("rating") if rating else "",
            "review_count": rating.get("review_count") if rating else "",
            "rating_provider": rating.get("provider") if rating else "",
            "rating_observed_at": rating.get("observed_at") if rating else "",
            "duration_short_minutes": profile.get("short_minutes", ""),
            "duration_typical_minutes": profile.get("typical_minutes", ""),
            "duration_long_minutes": profile.get("long_minutes", ""),
            "duration_method": profile.get("method", ""),
            "duration_verified_at": profile.get("verified_at") or "",
            "access_latitude": access.get("latitude", ""), "access_longitude": access.get("longitude", ""),
            "access_minutes": access.get("access_minutes", ""), "access_method": access.get("method", ""),
            "access_verified": bool(access.get("verified")), "entity_confirmed": bool(poi.get("entity_confirmed")),
            "serving_eligible": bool((poi.get("serving_quality") or {}).get("eligible")),
            "serving_reasons": "|".join((poi.get("serving_quality") or {}).get("reasons", [])),
            "data_rights": "mixed; see sources.csv and summary.json",
            "dataset_version": manifest["version"],
        })
    poi_fields = list(flat[0]) if flat else []
    write_csv(output_dir / "pois.csv", poi_fields, flat)

    with closing(connect(catalog, readonly=True)) as db:
        table_exports = {
            "opening_hours.csv": ("opening_intervals", ["interval_id", "poi_id", "day_of_week", "specific_date", "status", "open_minute", "close_minute", "closes_next_day", "observation_id"]),
            "ratings.csv": ("ratings", ["rating_id", "poi_id", "provider", "rating", "review_count", "observed_at", "same_observation", "observation_id"]),
            "duration_profiles.csv": ("duration_profiles", ["poi_id", "short_minutes", "typical_minutes", "long_minutes", "method", "observation_id", "verified_at"]),
            "access_points.csv": ("access_points", ["access_id", "poi_id", "latitude", "longitude", "kind", "access_minutes", "verified", "method", "observation_id"]),
            "sources.csv": ("source_files", ["source_file_id", "logical_path", "sha256", "role", "received_at", "tool", "use_status"]),
        }
        counts = {"pois.csv": len(flat)}
        for filename, (table, fields) in table_exports.items():
            rows = [dict(row) for row in db.execute(f"SELECT {','.join(fields)} FROM {table}")]
            write_csv(output_dir / filename, fields, rows)
            counts[filename] = len(rows)

    summary = dataset_summary(pois, manifest)
    summary["export_counts"] = counts
    if workbook_path and private_dir:
        summary["private_enrichment"] = export_private_enrichment(workbook_path, private_dir)
    (output_dir / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    readme = f"""# Dataset Where2Go DSS v2\n\nPhiên bản: `{manifest['version']}`  \nSinh lúc: `{manifest.get('created_at', '')}`\n\n- `pois.csv`: catalog phẳng, một dòng cho mỗi POI.\n- `opening_hours.csv`: các khoảng giờ cấu trúc; `day_of_week` dùng 0=Thứ Hai đến 6=Chủ Nhật.\n- `ratings.csv`: rating và số review theo cùng quan sát khi `same_observation=1`.\n- `duration_profiles.csv`: ba mức thời lượng; `category_default` chỉ là fallback thiết kế.\n- `access_points.csv`: tọa độ tiếp cận; `verified=0` nghĩa chưa được xác minh thủ công.\n- `sources.csv`: nguồn, checksum, vai trò và quyền sử dụng.\n- `summary.json`: coverage toàn catalog và tập ưu tiên 70 POI mỗi thành phố.\n\nCatalog chuẩn vẫn là `data/catalog_v2.sqlite`, chứa provenance chi tiết. Quan sát Google đang chờ kiểm duyệt được xuất riêng tại `data/private/google_enrichment_v2.csv` và `data/private/google_opening_hours_v2.csv`; hai file này bị Git bỏ qua. Dữ liệu Google Maps có trạng thái `restricted_internal`; không coi thư mục xuất này là dataset công khai được phép tái phân phối.\n"""
    readme = readme.replace("  \n", "\n\n")
    (output_dir / "README.md").write_text(readme, encoding="utf-8")
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--catalog", type=Path, default=CATALOG_V2)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "data/reports/v2/dataset")
    parser.add_argument("--workbook", type=Path, default=ROOT / "data/manual/poi_enrichment_v2.xlsx")
    parser.add_argument("--private-dir", type=Path, default=ROOT / "data/private")
    args = parser.parse_args()
    summary = export(args.catalog, args.output_dir, args.workbook, args.private_dir)
    print(json.dumps({"dataset_version": summary["dataset_version"], "poi_count": summary["poi_count"], "output_dir": str(args.output_dir)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
