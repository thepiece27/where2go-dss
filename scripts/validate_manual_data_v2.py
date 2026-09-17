"""Validate the manual enrichment workbook without mutating it."""
import argparse
from datetime import date, datetime, time
import json
import math
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from scripts.create_manual_template_v2 import PLACES, OPENING, DURATION, VERIFICATION
from where2go.config import ROOT


REQUIRED_SHEETS = {
    "places": PLACES,
    "opening_hours": OPENING,
    "visit_duration": DURATION,
    "verification": VERIFICATION,
}
DAYS = {"Mo", "Tu", "We", "Th", "Fr", "Sa", "Su"}


def blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def rows(sheet):
    headers = [cell.value for cell in sheet[1]]
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        item = dict(zip(headers, values))
        if any(not blank(value) for value in values):
            yield row_number, item


def validate(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    errors, warnings = [], []
    for sheet, expected in REQUIRED_SHEETS.items():
        if sheet not in workbook.sheetnames:
            errors.append({"sheet": sheet, "row": 1, "field": "sheet", "message": "Thiếu sheet bắt buộc"})
            continue
        actual = [cell.value for cell in workbook[sheet][1]]
        if actual != expected:
            errors.append({"sheet": sheet, "row": 1, "field": "headers", "message": "Header không đúng schema v2"})
    if errors:
        return errors, warnings, {}

    place_ids, canonical_ids = set(), set()
    for row_number, item in rows(workbook["places"]):
        ident = str(item.get("record_id") or "").strip()
        if not ident:
            errors.append({"sheet": "places", "row": row_number, "field": "record_id", "message": "Bắt buộc"})
        elif ident in place_ids:
            errors.append({"sheet": "places", "row": row_number, "field": "record_id", "message": "Bị lặp"})
        place_ids.add(ident)
        canonical = str(item.get("canonical_poi_id") or "").strip()
        if canonical:
            canonical_ids.add(canonical)
        location = item.get("location_expected")
        if location not in ("Hà Nội", "Đà Nẵng"):
            errors.append({"sheet": "places", "row": row_number, "field": "location_expected", "message": "Chỉ nhận Hà Nội/Đà Nẵng"})
        lat, lon = item.get("latitude"), item.get("longitude")
        if blank(lat) != blank(lon):
            errors.append({"sheet": "places", "row": row_number, "field": "coordinates", "message": "Vĩ độ/kinh độ phải cùng có hoặc cùng thiếu"})
        if not blank(lat):
            if not all(isinstance(v, (int, float)) and math.isfinite(v) for v in (lat, lon)) or not (-90 <= lat <= 90 and -180 <= lon <= 180):
                errors.append({"sheet": "places", "row": row_number, "field": "coordinates", "message": "Tọa độ không hợp lệ"})
            if blank(item.get("coordinate_method")):
                errors.append({"sheet": "places", "row": row_number, "field": "coordinate_method", "message": "Tọa độ cần phương pháp"})
        for field in ("rating_raw", "review_count_raw"):
            if item.get(field) == 0:
                warnings.append({"sheet": "places", "row": row_number, "field": field, "message": "0 có thể đang được dùng thay dữ liệu thiếu"})

    opening_keys = set()
    for row_number, item in rows(workbook["opening_hours"]):
        ident = str(item.get("record_id") or "").strip()
        if ident not in place_ids:
            errors.append({"sheet": "opening_hours", "row": row_number, "field": "record_id", "message": "Không tồn tại trong places"})
        day, specific, status = item.get("day_of_week"), item.get("specific_date"), item.get("status")
        if bool(day) == bool(specific):
            errors.append({"sheet": "opening_hours", "row": row_number, "field": "day/date", "message": "Chọn đúng một day_of_week hoặc specific_date"})
        if day and day not in DAYS:
            errors.append({"sheet": "opening_hours", "row": row_number, "field": "day_of_week", "message": "Mã ngày không hợp lệ"})
        if specific and not isinstance(specific, (date, datetime)):
            errors.append({"sheet": "opening_hours", "row": row_number, "field": "specific_date", "message": "Ngày không hợp lệ"})
        if status not in ("open", "closed", "unknown"):
            errors.append({"sheet": "opening_hours", "row": row_number, "field": "status", "message": "Trạng thái không hợp lệ"})
        start, end = item.get("open_time"), item.get("close_time")
        if status == "open" and not isinstance(start, time):
            errors.append({"sheet": "opening_hours", "row": row_number, "field": "open_time", "message": "open cần giờ mở"})
        if status == "open" and not isinstance(end, time):
            errors.append({"sheet": "opening_hours", "row": row_number, "field": "close_time", "message": "open cần giờ đóng"})
        if status in ("closed", "unknown") and (not blank(start) or not blank(end)):
            errors.append({"sheet": "opening_hours", "row": row_number, "field": "time", "message": "closed/unknown không được có khoảng giờ"})
        key = (ident, day, str(specific), status, str(start), str(end))
        if key in opening_keys:
            errors.append({"sheet": "opening_hours", "row": row_number, "field": "row", "message": "Khoảng giờ bị lặp"})
        opening_keys.add(key)

    for row_number, item in rows(workbook["visit_duration"]):
        ident = str(item.get("record_id") or "").strip()
        if ident not in place_ids:
            errors.append({"sheet": "visit_duration", "row": row_number, "field": "record_id", "message": "Không tồn tại trong places"})
        values = [item.get(name) for name in ("short_minutes", "typical_minutes", "long_minutes")]
        if not all(isinstance(value, int) and 5 <= value <= 720 for value in values):
            errors.append({"sheet": "visit_duration", "row": row_number, "field": "duration", "message": "Ba mức phải là số nguyên 5–720"})
        elif values != sorted(values):
            errors.append({"sheet": "visit_duration", "row": row_number, "field": "duration", "message": "short <= typical <= long"})
        if blank(item.get("duration_method")):
            errors.append({"sheet": "visit_duration", "row": row_number, "field": "duration_method", "message": "Bắt buộc"})

    verification_ids = set()
    for row_number, item in rows(workbook["verification"]):
        ident = str(item.get("record_id") or "").strip()
        if ident not in place_ids:
            errors.append({"sheet": "verification", "row": row_number, "field": "record_id", "message": "Không tồn tại trong places"})
        if ident in verification_ids:
            errors.append({"sheet": "verification", "row": row_number, "field": "record_id", "message": "Mỗi record chỉ có một kết luận hiện hành"})
        verification_ids.add(ident)
        if item.get("identity_status") not in ("confirmed", "ambiguous", "unmatched", "blocked", "closed"):
            errors.append({"sheet": "verification", "row": row_number, "field": "identity_status", "message": "Trạng thái không hợp lệ"})
        if item.get("identity_status") == "confirmed" and (blank(item.get("reviewer")) or blank(item.get("evidence_url"))):
            errors.append({"sheet": "verification", "row": row_number, "field": "evidence", "message": "confirmed cần reviewer và evidence URL"})

    stats = {
        "places": len(place_ids), "canonical_ids": len(canonical_ids),
        "opening_rows": len(opening_keys), "verification_rows": len(verification_ids),
    }
    return errors, warnings, stats


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", nargs="?", type=Path, default=ROOT / "data/manual/poi_enrichment_v2.xlsx")
    parser.add_argument("--report", type=Path, default=ROOT / "data/reports/v2/manual_validation.json")
    args = parser.parse_args()
    errors, warnings, stats = validate(args.workbook)
    payload = {"workbook": str(args.workbook), "status": "PASS" if not errors else "FAIL", "stats": stats, "errors": errors, "warnings": warnings}
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    print(json.dumps({"status": payload["status"], "errors": len(errors), "warnings": len(warnings), "report": str(args.report)}, ensure_ascii=False))
    raise SystemExit(1 if errors else 0)


if __name__ == "__main__":
    main()
