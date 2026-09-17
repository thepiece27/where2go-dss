"""Create the manual v2 enrichment workbook and a stratified 60-POI pilot seed."""
import argparse
from collections import defaultdict
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from where2go.catalog import load_catalog
from where2go.config import CATALOG, ROOT
from where2go.v2.catalog import CATALOG_V2, load_catalog as load_catalog_v2
from where2go.v2.dataset import select_priority
from where2go.v2.taxonomy import FOOD_CATEGORIES


PLACES = [
    "record_id", "canonical_poi_id", "seed_name", "maps_name", "maps_url", "place_id", "cid",
    "address_raw", "location_expected", "category_raw", "rating_raw", "review_count_raw",
    "latitude", "longitude", "coordinate_method", "website", "business_status_raw",
    "observed_at", "reviewer", "notes",
]
OPENING = [
    "record_id", "day_of_week", "specific_date", "status", "open_time", "close_time",
    "closes_next_day", "source_url", "observed_at", "notes",
]
DURATION = [
    "record_id", "short_minutes", "typical_minutes", "long_minutes", "duration_method",
    "source_url", "source_text_short", "observed_at", "reviewer", "notes",
]
VERIFICATION = [
    "record_id", "reviewer", "reviewed_at", "fields_checked", "identity_status",
    "evidence_url", "conflicts", "include_in_demo", "notes",
]


def _round_robin(rows, limit):
    groups = defaultdict(list)
    for poi in rows:
        groups[poi["category"]].append(poi)
    for items in groups.values():
        items.sort(key=lambda p: (
            not bool(p.get("ratings") or p.get("review")),
            p.get("hours_weekly") is None and not bool(p.get("hours_raw")), p["poi_id"],
        ))
    selected = []
    while len(selected) < limit:
        progressed = False
        for category in sorted(groups):
            if groups[category] and len(selected) < limit:
                selected.append(groups[category].pop(0))
                progressed = True
        if not progressed:
            break
    return selected


def select_seed(pois, per_city=30, food_per_city=10):
    """Round-robin categories to avoid a pilot made only of the largest class."""
    selected = []
    for location in ("Hà Nội", "Đà Nẵng"):
        if all("serving_quality" in poi for poi in pois):
            food_target = min(food_per_city, per_city // 3)
            selected.extend(select_priority(pois, location, False, per_city - food_target))
            selected.extend(select_priority(pois, location, True, food_target))
            continue
        eligible = [poi for poi in pois if poi.get("location") == location and poi.get("data_status") == "usable"]
        food = [poi for poi in eligible if poi.get("category") in FOOD_CATEGORIES]
        attractions = [poi for poi in eligible if poi.get("category") not in FOOD_CATEGORIES]
        food_target = min(food_per_city, len(food), per_city // 3)
        selected.extend(_round_robin(attractions, per_city - food_target))
        selected.extend(_round_robin(food, food_target))
    return selected


def style_sheet(sheet, widths):
    fill = PatternFill("solid", fgColor="155E75")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def create(output, catalog_path, per_city=30, force=False):
    if output.exists() and not force:
        raise FileExistsError(f"{output} đã tồn tại; dùng --force chỉ khi chưa có dữ liệu nhập tay cần giữ")
    if Path(catalog_path).name == "catalog_v2.sqlite":
        pois, _ = load_catalog_v2(catalog_path)
    else:
        pois, _ = load_catalog(catalog_path)
    seeds = select_seed(pois, per_city)
    workbook = Workbook()
    places = workbook.active
    places.title = "places"
    places.append(PLACES)
    for index, poi in enumerate(seeds, 1):
        places.append([
            f"pilot-{index:03d}", poi["poi_id"], poi["name"], "", "", "", "", "",
            poi["location"], "", "", "", "", "", "", poi.get("website", ""),
            "", "", "", "Đối soát pilot; không dùng tọa độ OSM làm dữ liệu Google nhập tay.",
        ])
    style_sheet(places, {"A": 16, "B": 26, "C": 34, "D": 34, "E": 48, "H": 34, "T": 48})

    opening = workbook.create_sheet("opening_hours")
    opening.append(OPENING)
    style_sheet(opening, {"A": 16, "B": 14, "C": 16, "D": 14, "H": 48, "J": 42})
    duration = workbook.create_sheet("visit_duration")
    duration.append(DURATION)
    style_sheet(duration, {"A": 16, "E": 28, "F": 48, "G": 55, "J": 42})
    verification = workbook.create_sheet("verification")
    verification.append(VERIFICATION)
    style_sheet(verification, {"A": 16, "D": 35, "E": 18, "F": 48, "G": 45, "I": 45})

    lists = workbook.create_sheet("_lists")
    lists.sheet_state = "hidden"
    values = {
        "A": ["Hà Nội", "Đà Nẵng"],
        "B": ["open", "temporarily_closed", "permanently_closed", "unknown"],
        "C": ["confirmed", "tool_confirmed", "ambiguous", "unmatched", "blocked", "closed"],
        "D": ["open", "closed", "unknown"],
        "E": ["TRUE", "FALSE"],
        "F": ["source_program", "reported_time_spent", "manual_estimate", "category_default"],
        "G": ["Mo", "Tu", "We", "Th", "Fr", "Sa", "Su"],
    }
    for column, entries in values.items():
        for row, value in enumerate(entries, 1):
            lists[f"{column}{row}"] = value

    validations = [
        (places, "I2:I1000", "'_lists'!$A$1:$A$2"),
        (places, "Q2:Q1000", "'_lists'!$B$1:$B$4"),
        (opening, "B2:B5000", "'_lists'!$G$1:$G$7"),
        (opening, "D2:D5000", "'_lists'!$D$1:$D$3"),
        (opening, "G2:G5000", "'_lists'!$E$1:$E$2"),
        (duration, "E2:E1000", "'_lists'!$F$1:$F$4"),
        (verification, "E2:E1000", "'_lists'!$C$1:$C$5"),
        (verification, "H2:H1000", "'_lists'!$E$1:$E$2"),
    ]
    for sheet, cells, formula in validations:
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        validation.error = "Giá trị không nằm trong danh sách cho phép"
        validation.errorTitle = "Dữ liệu không hợp lệ"
        validation.showErrorMessage = True
        sheet.add_data_validation(validation)
        validation.add(cells)

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return len(seeds)


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--catalog", type=Path, default=CATALOG_V2 if CATALOG_V2.exists() else CATALOG)
    parser.add_argument("--output", type=Path, default=ROOT / "data/manual/poi_enrichment_v2.xlsx")
    parser.add_argument("--per-city", type=int, default=30)
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args()
    count = create(args.output, args.catalog, args.per_city, args.force)
    print(f"Created {args.output} with {count} pilot seeds")


if __name__ == "__main__":
    main()
