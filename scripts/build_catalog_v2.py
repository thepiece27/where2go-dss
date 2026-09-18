"""Build the v2 observation database and conservative merged catalog."""
import argparse
import csv
from collections import Counter, defaultdict
from contextlib import closing
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook

from scripts.create_manual_template_v2 import PLACES, OPENING, DURATION, VERIFICATION
from scripts.inventory_sources_v2 import digest
from where2go.catalog import haversine, load_catalog
from where2go.config import CATALOG, ROOT
from where2go.hours import parse_week
from where2go.ranking import normalize
from where2go.v2 import MODEL_VERSION, SCHEMA_VERSION
from where2go.v2.durations import fallback_profile
from where2go.v2.observations import (
    parse_rating,
    parse_review_count, scalar, stable_id,
)
from where2go.v2.storage import connect, create_database, json_text
from where2go.v2.taxonomy import canonical_category, refine_category, tags_for


GOOGLE_FILES = (
    "vietnam_destinations_google_maps_browser_hotosm.xlsx",
    "vietnam_destinations_google_maps_browser_hotosm_backup_20260915_163012.xlsx",
    "vietnam_destinations_google_maps_browser_hotosm_backup_20260916_150600.xlsx",
)


def add_source(db, path, role, use_status, tool):
    sha = digest(path)
    ident = stable_id("source", path.relative_to(ROOT).as_posix(), sha)
    db.execute("INSERT INTO source_files VALUES (?,?,?,?,?,?,?)", (
        ident, path.relative_to(ROOT).as_posix(), sha, role, None, tool, use_status,
    ))
    return ident


def add_record(db, source_id, source_key, payload, observed_at):
    ident = stable_id("record", source_id, source_key)
    db.execute("INSERT INTO source_records VALUES (?,?,?,?,?)", (
        ident, source_id, str(source_key), observed_at, json_text(payload),
    ))
    return ident


def observe(db, poi_id, record_id, field, value, observed_at, use_status, method, verified_at=None, notes=None):
    ident = stable_id("obs", record_id, field)
    db.execute("INSERT OR REPLACE INTO field_observations VALUES (?,?,?,?,?,?,?,?,?,?)", (
        ident, poi_id, record_id, field, json_text(value), observed_at, verified_at,
        method, use_status, notes,
    ))
    return ident


def select(db, poi_id, field, observation_id, value, reason, build_version):
    db.execute("INSERT OR REPLACE INTO selected_fields VALUES (?,?,?,?,?,?)", (
        poi_id, field, observation_id, json_text(value), reason, build_version,
    ))


def import_base(db, catalog_path, build_version):
    pois, manifest = load_catalog(catalog_path)
    source_id = add_source(db, catalog_path, "generated_catalog", "open_data", "scripts/build_catalog.py")
    name_index = defaultdict(list)
    for poi in pois:
        effective_category = refine_category(poi["category"], poi["name"], poi.get("description", ""))
        record_id = add_record(db, source_id, poi["poi_id"], poi, manifest.get("created_at"))
        db.execute("INSERT INTO pois VALUES (?,?,?,?,?,?)", (
            poi["poi_id"], poi["data_status"], poi["name"], poi.get("location") or None,
            None, "unknown",
        ))
        db.execute("INSERT INTO source_links VALUES (?,?,?,?,?,?,?)", (
            poi["poi_id"], record_id, "v1_canonical_import", 1.0, None, "confirmed", None,
        ))
        fields = {
            "name": poi["name"], "latitude": poi["latitude"], "longitude": poi["longitude"],
            "location": poi.get("location"), "category": effective_category,
            "description": poi.get("description", ""), "hours_raw": poi.get("hours_raw", ""),
            "website": poi.get("website", ""), "source_url": poi.get("source_url", ""),
            "review_reasons": poi.get("review_reason", []),
        }
        observations = {}
        for field, value in fields.items():
            provenance = poi.get("provenance", {}).get(field, {})
            observations[field] = observe(
                db, poi["poi_id"], record_id, field, value,
                provenance.get("observed_at") or manifest.get("created_at"),
                "open_data", provenance.get("method", "v1_catalog_import"),
                provenance.get("verified_at"),
            )
            select(db, poi["poi_id"], field, observations[field], value, "v1 canonical baseline", build_version)
        db.execute("INSERT INTO poi_categories VALUES (?,?,1,?)", (poi["poi_id"], effective_category, observations["category"]))
        for tag in tags_for(effective_category):
            db.execute("INSERT INTO poi_categories VALUES (?,?,0,?)", (poi["poi_id"], "tag:" + tag, observations["category"]))
        profile = fallback_profile(effective_category)
        db.execute("INSERT INTO duration_profiles VALUES (?,?,?,?,?,?,?)", (
            poi["poi_id"], profile["short_minutes"], profile["typical_minutes"], profile["long_minutes"],
            profile["method"], None, None,
        ))
        db.execute("INSERT INTO access_points VALUES (?,?,?,?,?,?,?,?,?)", (
            stable_id("access", poi["poi_id"], "base"), poi["poi_id"], poi["latitude"], poi["longitude"],
            "poi_coordinate", 10, 0, poi.get("coordinate_status", "v1_coordinate"), observations["latitude"],
        ))
        weekly = parse_week(poi.get("hours_raw"))
        if weekly is not None:
            for weekday, intervals in enumerate(weekly):
                if intervals is None:
                    continue
                if not intervals:
                    db.execute("INSERT INTO opening_intervals VALUES (?,?,?,?,?,?,?,?,?)", (
                        stable_id("hours", poi["poi_id"], weekday, "closed"), poi["poi_id"], observations["hours_raw"],
                        weekday, None, "closed", None, None, 0,
                    ))
                for start, end in intervals:
                    db.execute("INSERT INTO opening_intervals VALUES (?,?,?,?,?,?,?,?,?)", (
                        stable_id("hours", poi["poi_id"], weekday, start, end), poi["poi_id"], observations["hours_raw"],
                        weekday, None, "open", start, end if end <= 1440 else end - 1440, int(end > 1440),
                    ))
        name_index[(normalize(poi["name"]), poi.get("location"))].append(dict(poi, category=effective_category))
    return pois, manifest, name_index


def import_osm_supplement(db, pbf_path, supplement_path, name_index, build_version):
    if not supplement_path.exists():
        raise FileNotFoundError(f"Thiếu {supplement_path}; chạy scripts/extract_osm_supplement_v2.py")
    payload = json.loads(supplement_path.read_text(encoding="utf-8"))
    source_id = add_source(db, pbf_path, "spatial_source", "open_data", "scripts/extract_osm_supplement_v2.py")
    counters = Counter()
    for row in payload["rows"]:
        record_id = add_record(db, source_id, row["poi_id"], row, row.get("observed_at"))
        peers = name_index.get((row["name_normalized"], row["location"]), [])
        duplicate = next((p for p in peers if haversine((row["latitude"], row["longitude"]),
                                                         (p["latitude"], p["longitude"])) < .15), None)
        if duplicate:
            db.execute("INSERT INTO source_links VALUES (?,?,?,?,?,?,?)", (
                duplicate["poi_id"], record_id, "exact_name_location_within_150m", 1.0,
                None, "confirmed", "supplement linked to existing canonical POI",
            ))
            counters["linked_existing"] += 1
            continue
        poi_id = row["poi_id"]
        db.execute("INSERT INTO pois VALUES (?,?,?,?,?,?)", (
            poi_id, "usable", row["name"], row["location"], None, "unknown",
        ))
        db.execute("INSERT INTO source_links VALUES (?,?,?,?,?,?,?)", (
            poi_id, record_id, "osm_source_id", 1.0, None, "confirmed", None,
        ))
        observations = {}
        for field in ("name", "latitude", "longitude", "location", "category", "description", "hours_raw", "website", "source_url"):
            value = row.get(field)
            observations[field] = observe(
                db, poi_id, record_id, field, value, row.get("observed_at"), "open_data",
                row.get("coordinate_method") if field in ("latitude", "longitude") else "osm_tag",
            )
            select(db, poi_id, field, observations[field], value, "OSM supplemental observation", build_version)
        db.execute("INSERT INTO poi_categories VALUES (?,?,1,?)", (poi_id, row["category"], observations["category"]))
        for tag in tags_for(row["category"]):
            db.execute("INSERT INTO poi_categories VALUES (?,?,0,?)", (poi_id, "tag:" + tag, observations["category"]))
        profile = fallback_profile(row["category"])
        db.execute("INSERT INTO duration_profiles VALUES (?,?,?,?,?,?,?)", (
            poi_id, profile["short_minutes"], profile["typical_minutes"], profile["long_minutes"],
            profile["method"], None, None,
        ))
        db.execute("INSERT INTO access_points VALUES (?,?,?,?,?,?,?,?,?)", (
            stable_id("access", poi_id, "base"), poi_id, row["latitude"], row["longitude"],
            "poi_coordinate", 10, 0, row["coordinate_method"], observations["latitude"],
        ))
        weekly = parse_week(row.get("hours_raw"))
        if weekly is not None:
            for weekday, intervals in enumerate(weekly):
                if intervals is None:
                    continue
                if not intervals:
                    db.execute("INSERT INTO opening_intervals VALUES (?,?,?,?,?,?,?,?,?)", (
                        stable_id("hours", poi_id, weekday, "closed"), poi_id, observations["hours_raw"],
                        weekday, None, "closed", None, None, 0,
                    ))
                for start, end in intervals:
                    db.execute("INSERT INTO opening_intervals VALUES (?,?,?,?,?,?,?,?,?)", (
                        stable_id("hours", poi_id, weekday, start, end), poi_id, observations["hours_raw"],
                        weekday, None, "open", start, end if end <= 1440 else end - 1440, int(end > 1440),
                    ))
        canonical = {
            "poi_id": poi_id, "name": row["name"], "name_normalized": row["name_normalized"],
            "latitude": row["latitude"], "longitude": row["longitude"], "location": row["location"],
            "category": row["category"], "data_status": "usable",
        }
        name_index[(row["name_normalized"], row["location"])].append(canonical)
        counters[row["location"] + ":" + row["category"]] += 1
        counters["inserted"] += 1
    return counters


def import_manual(db, path, queue, build_version):
    if not path.exists():
        return Counter(missing=1)
    workbook = load_workbook(path, read_only=True, data_only=True)
    expected = {
        "places": PLACES, "opening_hours": OPENING,
        "visit_duration": DURATION, "verification": VERIFICATION,
    }
    for sheet_name, columns in expected.items():
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Manual workbook is missing sheet {sheet_name}")
        headers = [cell.value for cell in workbook[sheet_name][1]]
        if headers != columns:
            raise ValueError(f"Manual workbook headers do not match v2 schema: {sheet_name}")

    def sheet_rows(sheet_name, columns):
        for row_number, values in enumerate(
                workbook[sheet_name].iter_rows(min_row=2, values_only=True), 2):
            if not any(value not in (None, "") for value in values):
                continue
            yield row_number, dict(zip(columns, values))

    verification = {}
    for _, raw in sheet_rows("verification", VERIFICATION):
        row = {key: scalar(value) for key, value in raw.items()}
        verification[str(row.get("record_id") or "").strip()] = row

    source_id = add_source(db, path, "manual_observation", "manual_fact_with_source", "manual workbook")
    counters = Counter()
    records = {}
    confirmed = {}
    for row_number, raw in sheet_rows("places", PLACES):
        row = {key: scalar(value) for key, value in raw.items()}
        record_id = add_record(db, source_id, row["record_id"], row, row.get("observed_at"))
        records[str(row["record_id"])] = record_id
        poi_id = str(row.get("canonical_poi_id") or "").strip()
        exists = poi_id and db.execute("SELECT 1 FROM pois WHERE poi_id=?", (poi_id,)).fetchone()
        decision = verification.get(str(row["record_id"]), {})
        identity_status = decision.get("identity_status")
        created_new = False
        identity_confirmed = identity_status in ("confirmed", "tool_confirmed")
        if not exists and identity_confirmed:
            latitude, longitude = row.get("latitude"), row.get("longitude")
            category = canonical_category(row.get("category_raw"), row.get("maps_name") or row.get("seed_name"))
            identity_key = row.get("place_id") or row.get("cid") or row.get("maps_url")
            complete = (
                row.get("maps_name") and row.get("maps_url") and identity_key and category
                and isinstance(latitude, (int, float)) and isinstance(longitude, (int, float))
            )
            if complete:
                poi_id = stable_id("poi-google", identity_key)
                business = str(row.get("business_status_raw") or "unknown").strip()
                if business not in ("open", "temporarily_closed", "permanently_closed", "unknown"):
                    business = "unknown"
                db.execute("INSERT INTO pois VALUES (?,?,?,?,?,?)", (
                    poi_id, "usable", row["maps_name"], row.get("location_expected"), None, business,
                ))
                exists = True
                created_new = True
                counters["new_pois"] += 1
            else:
                queue.append({
                    "source_file": path.name, "source_key": row["record_id"], "status": "ambiguous",
                    "candidate_poi_ids": [],
                    "reasons": ["confirmed_new_poi_requires_name_url_category_coordinates"],
                })
                counters["incomplete_new_poi"] += 1
        if exists and identity_confirmed:
            db.execute("INSERT INTO source_links VALUES (?,?,?,?,?,?,?)", (
                poi_id, record_id,
                "tool_cross_source_identity_verification" if identity_status == "tool_confirmed" else "manual_identity_verification",
                1.0,
                decision.get("reviewer") or row.get("reviewer"), "confirmed",
                decision.get("notes") or "Identity confirmed in verification sheet",
            ))
            confirmed[str(row["record_id"])] = (poi_id, row, decision)
            counters["confirmed"] += 1
            if identity_status == "tool_confirmed":
                counters["tool_confirmed"] += 1
            observed_at = row.get("observed_at")
            verified_at = decision.get("reviewed_at") if identity_status == "confirmed" else None
            fields = {
                "google_maps_url": row.get("maps_url"), "name": row.get("maps_name"),
                "website": row.get("website"), "address_raw": row.get("address_raw"),
                "google_category_raw": row.get("category_raw"),
                "google_place_id": row.get("place_id"), "google_cid": row.get("cid"),
            }
            observations = {}
            for field, value in fields.items():
                if value in (None, ""):
                    continue
                observation = observe(
                    db, poi_id, record_id, field, value, observed_at,
                    "manual_fact_with_source", "manual_verified_observation",
                    verified_at, decision.get("evidence_url"),
                )
                observations[field] = observation
                if field == "google_place_id":
                    db.execute("INSERT OR IGNORE INTO external_ids VALUES ('Google Maps',?,?,?)", (value, poi_id, observation))
                if field == "name":
                    db.execute("INSERT OR IGNORE INTO poi_aliases VALUES (?,?,?,?)", (poi_id, value, normalize(value), observation))
                if field == "website" or (created_new and field == "name"):
                    select(db, poi_id, field, observation, value,
                           "confirmed manual observation", build_version)

            business = str(row.get("business_status_raw") or "").strip()
            if business in ("open", "temporarily_closed", "permanently_closed", "unknown"):
                db.execute("UPDATE pois SET business_status=? WHERE poi_id=?", (business, poi_id))

            category = canonical_category(row.get("category_raw"), row.get("maps_name"))
            if created_new:
                latitude, longitude = row["latitude"], row["longitude"]
                for field, value in (("latitude", latitude), ("longitude", longitude),
                                     ("location", row.get("location_expected")), ("category", category),
                                     ("source_url", row.get("maps_url"))):
                    observation = observe(
                        db, poi_id, record_id, field, value, observed_at,
                        "restricted_internal" if field != "location" else "manual_fact_with_source",
                        "manual_verified_observation", verified_at, decision.get("evidence_url"),
                    )
                    observations[field] = observation
                    select(db, poi_id, field, observation, value, "confirmed new manual POI", build_version)
                db.execute("INSERT INTO poi_categories VALUES (?,?,1,?)", (
                    poi_id, category, observations["category"],
                ))
                for tag in tags_for(category):
                    db.execute("INSERT INTO poi_categories VALUES (?,?,0,?)", (
                        poi_id, "tag:" + tag, observations["category"],
                    ))
                profile = fallback_profile(category)
                db.execute("INSERT INTO duration_profiles VALUES (?,?,?,?,?,?,?)", (
                    poi_id, profile["short_minutes"], profile["typical_minutes"], profile["long_minutes"],
                    profile["method"], None, None,
                ))
            elif category:
                current_category = db.execute(
                    "SELECT category FROM poi_categories WHERE poi_id=? AND is_primary=1", (poi_id,)
                ).fetchone()
                if current_category and current_category[0] == "attraction" and category != "attraction":
                    observation = observe(
                        db, poi_id, record_id, "category", category, observed_at,
                        "manual_fact_with_source", "manual_verified_google_category",
                        verified_at, decision.get("evidence_url"),
                    )
                    db.execute("DELETE FROM poi_categories WHERE poi_id=? AND is_primary=1", (poi_id,))
                    db.execute("INSERT OR IGNORE INTO poi_categories VALUES (?,?,1,?)", (poi_id, category, observation))
                    select(db, poi_id, "category", observation, category,
                           "refined broad category from confirmed observation", build_version)

            rating = parse_rating(row.get("rating_raw"))
            count = parse_review_count(row.get("review_count_raw"))
            if rating is not None and not math.isclose(rating * 10, round(rating * 10), abs_tol=1e-8):
                rating = None
                counters["invalid_google_rating_precision"] += 1
            if rating is not None or count is not None:
                observation = observe(
                    db, poi_id, record_id, "rating_pair",
                    {"rating": rating, "review_count": count, "provider": "Google Maps"},
                    observed_at, "restricted_internal", "manual_same_source_page",
                    verified_at, decision.get("evidence_url"),
                )
                db.execute("INSERT OR REPLACE INTO ratings VALUES (?,?,?,?,?,?,?,?)", (
                    stable_id("rating", record_id), poi_id, "Google Maps", rating, count,
                    observed_at, int(rating is not None and count is not None), observation,
                ))

            lat, lon = row.get("latitude"), row.get("longitude")
            if isinstance(lat, (int, float)) and isinstance(lon, (int, float)):
                method = row.get("coordinate_method") or "manual_map_selection"
                access_verified = str(method).lower() in (
                    "verified_entrance", "verified_access", "manual_verified_entrance",
                )
                observation = observe(
                    db, poi_id, record_id, "access_coordinate", {"latitude": lat, "longitude": lon},
                    observed_at, "manual_fact_with_source", method,
                    verified_at, decision.get("evidence_url"),
                )
                db.execute("INSERT OR REPLACE INTO access_points VALUES (?,?,?,?,?,?,?,?,?)", (
                    stable_id("access", poi_id, "manual", row["record_id"]), poi_id, lat, lon,
                    "verified_entrance" if access_verified else "poi_coordinate", 0,
                    int(access_verified), method,
                    observation,
                ))
        else:
            if exists and (row.get("maps_name") or row.get("maps_url")):
                db.execute("INSERT INTO source_links VALUES (?,?,?,?,?,?,?)", (
                    poi_id, record_id, "manual_workbook_pending_verification", None,
                    row.get("reviewer"), "ambiguous", "Verification sheet must confirm identity",
                ))
                counters["pending_verification"] += 1
            else:
                counters["seed_only"] += 1
            if row.get("maps_name") or row.get("maps_url"):
                queue.append({"source_file": path.name, "source_key": row["record_id"], "status": "ambiguous",
                              "candidate_poi_ids": [poi_id] if exists else [],
                              "reasons": ["manual_verification_required"]})

    day_numbers = {"Mo": 0, "Tu": 1, "We": 2, "Th": 3, "Fr": 4, "Sa": 5, "Su": 6}
    cleared_hours = set()
    for _, raw in sheet_rows("opening_hours", OPENING):
        key = str(raw.get("record_id") or "").strip()
        if key not in confirmed:
            counters["hours_skipped_unconfirmed"] += 1
            continue
        poi_id, place, decision = confirmed[key]
        if poi_id not in cleared_hours:
            db.execute("DELETE FROM opening_intervals WHERE poi_id=?", (poi_id,))
            cleared_hours.add(poi_id)
            counters["hours_sources_replaced"] += 1
        record_id = records[key]
        status = str(raw.get("status") or "").strip()
        observed_at = scalar(raw.get("observed_at")) or place.get("observed_at")
        payload = {field: scalar(value) for field, value in raw.items()}
        observation = observe(
            db, poi_id, record_id, "opening_hours", payload, observed_at,
            "manual_fact_with_source", "manual_structured_hours",
            decision.get("reviewed_at"), scalar(raw.get("source_url")),
        )
        day = day_numbers.get(raw.get("day_of_week"))
        specific = raw.get("specific_date")
        specific_date = scalar(specific) if specific else None
        opens = raw.get("open_time")
        closes = raw.get("close_time")
        open_minute = opens.hour * 60 + opens.minute if hasattr(opens, "hour") else None
        close_minute = closes.hour * 60 + closes.minute if hasattr(closes, "hour") else None
        next_day = str(raw.get("closes_next_day") or "").strip().lower() in ("true", "1", "yes")
        db.execute("INSERT OR REPLACE INTO opening_intervals VALUES (?,?,?,?,?,?,?,?,?)", (
            stable_id("hours", record_id, day, specific_date, status, open_minute, close_minute, next_day),
            poi_id, observation, day, specific_date, status,
            open_minute if status == "open" else None,
            close_minute if status == "open" else None, int(next_day),
        ))
        counters["opening_rows"] += 1

    for _, raw in sheet_rows("visit_duration", DURATION):
        key = str(raw.get("record_id") or "").strip()
        if key not in confirmed:
            counters["duration_skipped_unconfirmed"] += 1
            continue
        poi_id, place, decision = confirmed[key]
        record_id = records[key]
        observed_at = scalar(raw.get("observed_at")) or place.get("observed_at")
        payload = {field: scalar(value) for field, value in raw.items()}
        observation = observe(
            db, poi_id, record_id, "visit_duration", payload, observed_at,
            "manual_fact_with_source", raw.get("duration_method"),
            decision.get("reviewed_at"), scalar(raw.get("source_url")),
        )
        db.execute("INSERT OR REPLACE INTO duration_profiles VALUES (?,?,?,?,?,?,?)", (
            poi_id, int(raw["short_minutes"]), int(raw["typical_minutes"]),
            int(raw["long_minutes"]), str(raw["duration_method"]), observation,
            decision.get("reviewed_at"),
        ))
        counters["duration_profiles"] += 1
    return counters


def import_duration_curation(db, path, build_version):
    counters = Counter()
    if not path.exists():
        return counters
    source_id = add_source(
        db, path, "planning_assumption", "internal_design_data",
        "data/curation/focus_duration_profiles_v2.csv",
    )
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    for row in rows:
        poi_id = str(row.get("poi_id") or "").strip()
        if not db.execute("SELECT 1 FROM pois WHERE poi_id=?", (poi_id,)).fetchone():
            counters["missing_poi"] += 1
            continue
        try:
            short = int(row["short_minutes"])
            typical = int(row["typical_minutes"])
            long = int(row["long_minutes"])
        except (TypeError, ValueError):
            counters["invalid"] += 1
            continue
        if not (0 < short <= typical <= long <= 1440):
            counters["invalid"] += 1
            continue
        record_id = add_record(db, source_id, poi_id, row, None)
        payload = {
            "short_minutes": short, "typical_minutes": typical, "long_minutes": long,
            "method": row.get("method") or "curated_planning_estimate",
            "source_url": row.get("source_url") or None, "notes": row.get("notes") or None,
        }
        observation = observe(
            db, poi_id, record_id, "visit_duration", payload, None,
            "internal_design_data", payload["method"], None, payload["source_url"],
        )
        db.execute("INSERT OR REPLACE INTO duration_profiles VALUES (?,?,?,?,?,?,?)", (
            poi_id, short, typical, long, payload["method"], observation, None,
        ))
        counters["imported"] += 1
    return counters


def import_poi_relations(db, path, build_version):
    counters = Counter()
    if not path.exists():
        return counters
    source_id = add_source(
        db, path, "entity_relation_curation", "open_data_derived",
        "data/curation/poi_relations_v2.csv",
    )
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    for row in rows:
        child = str(row.get("child_poi_id") or "").strip()
        parent = str(row.get("parent_poi_id") or "").strip()
        if child == parent or not db.execute("SELECT 1 FROM pois WHERE poi_id=?", (child,)).fetchone() \
                or not db.execute("SELECT 1 FROM pois WHERE poi_id=?", (parent,)).fetchone():
            counters["invalid_or_missing"] += 1
            continue
        record_id = add_record(db, source_id, child, row, None)
        observation = observe(
            db, child, record_id, "parent_poi_id", parent, None,
            "open_data_derived", row.get("relation") or "contained_experience",
            None, row.get("notes"),
        )
        db.execute("UPDATE pois SET parent_poi_id=? WHERE poi_id=?", (parent, child))
        select(db, child, "parent_poi_id", observation, parent, "curated contained experience", build_version)
        counters["imported"] += 1
    return counters


def import_focus_category_curation(db, path, build_version):
    counters = Counter()
    if not path.exists():
        return counters
    source_id = add_source(
        db, path, "focus_taxonomy_curation", "internal_design_data",
        "data/curation/focus_landmarks_v2.csv",
    )
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    for row in rows:
        poi_id = str(row.get("canonical_poi_id") or "").strip()
        category = str(row.get("expected_category") or "").strip()
        if not poi_id or not category or not db.execute("SELECT 1 FROM pois WHERE poi_id=?", (poi_id,)).fetchone():
            counters["missing_or_new_poi"] += 1
            continue
        current = db.execute(
            "SELECT category FROM poi_categories WHERE poi_id=? AND is_primary=1", (poi_id,)
        ).fetchone()
        poi_row = db.execute("SELECT location FROM pois WHERE poi_id=?", (poi_id,)).fetchone()
        expected_location = str(row.get("location_expected") or "").strip()
        update_location = bool(expected_location and not poi_row[0])
        update_category = not current or current[0] != category
        if not update_category and not update_location:
            counters["unchanged"] += 1
            continue
        record_id = add_record(db, source_id, poi_id, row, None)
        if update_category:
            observation = observe(
                db, poi_id, record_id, "category", category, None,
                "internal_design_data", "curated_focus_taxonomy", None, row.get("reason"),
            )
            db.execute("DELETE FROM poi_categories WHERE poi_id=?", (poi_id,))
            db.execute("INSERT OR REPLACE INTO poi_categories VALUES (?,?,1,?)", (poi_id, category, observation))
            for tag in tags_for(category):
                db.execute("INSERT OR IGNORE INTO poi_categories VALUES (?,?,0,?)", (
                    poi_id, "tag:" + tag, observation,
                ))
            select(db, poi_id, "category", observation, category, "curated focus taxonomy", build_version)
            counters["category_updated"] += 1
        if update_location:
            observation = observe(
                db, poi_id, record_id, "location", expected_location, None,
                "open_data_derived", "curated_spatial_assignment", None, row.get("reason"),
            )
            db.execute("UPDATE pois SET location=? WHERE poi_id=?", (expected_location, poi_id))
            select(db, poi_id, "location", observation, expected_location,
                   "curated focus spatial assignment", build_version)
            counters["location_updated"] += 1
    return counters


def coverage(db):
    result = []
    for location in ("Hà Nội", "Đà Nẵng"):
        total = db.execute("SELECT count(*) FROM pois WHERE location=?", (location,)).fetchone()[0]
        usable = db.execute("SELECT count(*) FROM pois WHERE location=? AND status='usable'", (location,)).fetchone()[0]
        ratings = db.execute("SELECT count(DISTINCT r.poi_id) FROM ratings r JOIN pois p ON p.poi_id=r.poi_id WHERE p.location=? AND r.same_observation=1", (location,)).fetchone()[0]
        hours = db.execute("SELECT count(DISTINCT h.poi_id) FROM opening_intervals h JOIN pois p ON p.poi_id=h.poi_id WHERE p.location=?", (location,)).fetchone()[0]
        food = db.execute("SELECT count(DISTINCT p.poi_id) FROM pois p JOIN poi_categories c ON c.poi_id=p.poi_id AND c.is_primary=1 WHERE p.location=? AND p.status='usable' AND c.category IN ('restaurant','cafe','food_street')", (location,)).fetchone()[0]
        attractions = usable - food
        durations = db.execute("SELECT count(DISTINCT d.poi_id) FROM duration_profiles d JOIN pois p ON p.poi_id=d.poi_id WHERE p.location=? AND d.method<>'category_default'", (location,)).fetchone()[0]
        verified_durations = db.execute("SELECT count(DISTINCT d.poi_id) FROM duration_profiles d JOIN pois p ON p.poi_id=d.poi_id WHERE p.location=? AND d.verified_at IS NOT NULL", (location,)).fetchone()[0]
        curated_durations = db.execute("SELECT count(DISTINCT d.poi_id) FROM duration_profiles d JOIN pois p ON p.poi_id=d.poi_id WHERE p.location=? AND d.method='curated_planning_estimate'", (location,)).fetchone()[0]
        verified_access = db.execute("SELECT count(DISTINCT a.poi_id) FROM access_points a JOIN pois p ON p.poi_id=a.poi_id WHERE p.location=? AND a.verified=1", (location,)).fetchone()[0]
        result.append({"location": location, "total": total, "usable": usable,
                       "attractions": attractions, "food_rest": food, "with_rating_pair": ratings,
                       "with_structured_hours": hours, "with_specific_duration": durations,
                       "with_verified_duration": verified_durations,
                       "with_curated_duration_estimate": curated_durations,
                       "with_verified_access": verified_access})
    return result


def pipeline_hash(paths):
    hasher = hashlib.sha256()
    for path in sorted(paths, key=lambda item: str(item)):
        hasher.update(str(path.relative_to(ROOT)).encode())
        hasher.update(digest(path).encode())
    return hasher.hexdigest()


def build(args):
    from scripts.merge_sources_v2 import import_legacy, import_collected, apply_images, reconcile_status
    geometry_path = ROOT / "data/cache/poi_geometries_v2.json"
    image_checks = ROOT / "data/cache/image_checks_v2.json"
    collected_paths = sorted((ROOT / "data/enrichment").glob("accepted*.json"))
    input_paths = [args.v1_catalog, args.pbf, args.supplement, args.duration_curation,
                   args.poi_relations, args.focus_curation, *args.google]
    if args.manual.exists():
        input_paths.append(args.manual)
    input_paths.extend(path for path in [geometry_path, image_checks, *collected_paths] if path.exists())
    code_paths = [
        Path(__file__),
        ROOT / "where2go/hours.py",
        ROOT / "where2go/v2/storage.py",
        ROOT / "where2go/v2/__init__.py",
        ROOT / "where2go/v2/observations.py",
        ROOT / "where2go/v2/taxonomy.py",
        ROOT / "where2go/v2/durations.py",
        ROOT / "where2go/v2/google_hours.py",
        ROOT / "where2go/v2/catalog.py",
        ROOT / "where2go/v2/quality.py",
        ROOT / "where2go/v2/dataset.py",
        ROOT / "scripts/export_dataset_v2.py",
        ROOT / "scripts/merge_sources_v2.py",
        ROOT / "where2go/v2/google_collector.py",
    ]
    input_hash = pipeline_hash(input_paths + code_paths)
    build_version = "v2-" + input_hash[:16]
    staging_output = args.output.with_suffix(".building.sqlite")
    create_database(staging_output)
    review_queue = []
    with closing(connect(staging_output)) as db:
        pois, v1_manifest, name_index = import_base(db, args.v1_catalog, build_version)
        supplement_stats = import_osm_supplement(db, args.pbf, args.supplement, name_index, build_version)
        manual_stats = import_manual(db, args.manual, review_queue, build_version)
        duration_stats = import_duration_curation(db, args.duration_curation, build_version)
        relation_stats = import_poi_relations(db, args.poi_relations, build_version)
        category_stats = import_focus_category_curation(db, args.focus_curation, build_version)
        geometry_cache = json.loads(geometry_path.read_text(encoding="utf-8")) if geometry_path.exists() else {}
        if geometry_cache and geometry_cache.get("pbf_sha256") != v1_manifest["osm"]["sha256"]:
            raise ValueError("Geometry cache does not match the OSM snapshot")
        collected_stats = import_collected(db, collected_paths, review_queue, build_version)
        google_stats = import_legacy(db, args.google, review_queue, build_version, geometry_cache.get("geometries"))
        image_stats = apply_images(db, image_checks, build_version)
        resolved_statuses = reconcile_status(db)
        stats = {
            "poi_count": db.execute("SELECT count(*) FROM pois").fetchone()[0],
            "osm_supplement": dict(supplement_stats), "google": dict(google_stats), "manual": dict(manual_stats),
            "duration_curation": dict(duration_stats),
            "poi_relations": dict(relation_stats),
            "focus_category_curation": dict(category_stats),
            "collected": collected_stats, "images": image_stats, "resolved_spatial_statuses": resolved_statuses,
            "review_queue": len(review_queue), "focus_coverage": coverage(db),
        }
        db.execute("INSERT INTO builds VALUES (?,?,?,?,?)", (
            build_version, datetime.now(timezone.utc).isoformat(), input_hash, MODEL_VERSION, json_text(stats),
        ))
        manifest = {
            "version": build_version, "created_at": datetime.now(timezone.utc).isoformat(),
            "schema_version": SCHEMA_VERSION, "model_version": MODEL_VERSION, "input_hash": input_hash,
            "v1_dataset_version": v1_manifest["version"], "osm": v1_manifest["osm"], "stats": stats,
            "rights": {"Google Maps": "restricted_internal", "OpenStreetMap": "ODbL-1.0"},
        }
        db.execute("INSERT INTO metadata VALUES ('manifest',?)", (json_text(manifest),))
        db.commit()
        foreign = db.execute("PRAGMA foreign_key_check").fetchall()
        if foreign:
            raise RuntimeError(f"Foreign key errors: {foreign[:5]}")
    staging_output.replace(args.output)
    args.report_dir.mkdir(parents=True, exist_ok=True)
    (args.report_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    (args.report_dir / "coverage.json").write_text(json.dumps(stats["focus_coverage"], ensure_ascii=False, indent=2), encoding="utf-8")
    (args.report_dir / "review_queue.json").write_text(json.dumps(review_queue, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"version": build_version, **stats}, ensure_ascii=False))


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--v1-catalog", type=Path, default=CATALOG)
    parser.add_argument("--pbf", type=Path, default=ROOT / "data/raw/vietnam-260915.osm.pbf")
    parser.add_argument("--supplement", type=Path, default=ROOT / "data/cache/osm_supplement_v2.json")
    parser.add_argument("--google", type=Path, nargs="*", default=[ROOT / "data" / name for name in GOOGLE_FILES])
    parser.add_argument("--manual", type=Path, default=ROOT / "data/manual/poi_enrichment_v2.xlsx")
    parser.add_argument("--duration-curation", type=Path, default=ROOT / "data/curation/focus_duration_profiles_v2.csv")
    parser.add_argument("--poi-relations", type=Path, default=ROOT / "data/curation/poi_relations_v2.csv")
    parser.add_argument("--focus-curation", type=Path, default=ROOT / "data/curation/focus_landmarks_v2.csv")
    parser.add_argument("--output", type=Path, default=ROOT / "data/catalog_v2.sqlite")
    parser.add_argument("--report-dir", type=Path, default=ROOT / "data/reports/v2/catalog")
    build(parser.parse_args())


if __name__ == "__main__":
    main()
